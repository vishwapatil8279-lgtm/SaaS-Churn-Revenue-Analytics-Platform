import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

customers = pd.read_csv("cleaned_saas_customers.csv")
subscriptions = pd.read_csv("cleaned_saas_subscriptions.csv")
usage = pd.read_csv("cleaned_saas_usage.csv")
tickets = pd.read_csv("cleaned_saas_tickets.csv")

print("Customers:", customers.shape)
print("Subscriptions:", subscriptions.shape)
print("Usage:", usage.shape)
print("Tickets:", tickets.shape)

# STEP 2 — CONVERT DATES

customers["SignupDate"] = pd.to_datetime(
    customers["SignupDate"],
    errors="coerce"
)
subscriptions["StartDate"] = pd.to_datetime(
    subscriptions["StartDate"],
    errors="coerce"
)
usage["Month"] = pd.to_datetime(
    usage["Month"],
    errors="coerce"
)
tickets["OpenedDate"] = pd.to_datetime(
    tickets["OpenedDate"],
    errors="coerce"
)

# STEP 3 — CUSTOMER-LEVEL SUBSCRIPTION SUMMARY

subscription_summary = (
    subscriptions
    .groupby("CustomerID")
    .agg(
        TotalMRR=("MRR", "sum"),
        TotalSeats=("Seats", "sum"),
        SubscriptionCount=("CustomerID", "count")
    )
    .reset_index()
)

# STEP 4 — LATEST SUBSCRIPTION STATUS
latest_subscription = (
    subscriptions
    .sort_values(
        ["CustomerID", "StartDate"]
    )
    .groupby("CustomerID")
    .tail(1)
    .copy()
)
latest_status = latest_subscription[
    ["CustomerID", "Status"]
].copy()

# STEP 5 — CREATE CUSTOMER SUMMARY
customer_summary = customers[
    [
        "CustomerID",
        "CompanyName",
        "Industry",
        "Country",
        "City",
        "EmployeeCount",
        "SignupDate",
        "AcquisitionChannel"
    ]
].copy()

customer_summary = customer_summary.merge(
    subscription_summary,
    on="CustomerID",
    how="left"
)

customer_summary = customer_summary.merge(
    latest_status,
    on="CustomerID",
    how="left"
)
# STEP 6 — HANDLE MISSING VALUES

customer_summary["TotalMRR"] = (
    customer_summary["TotalMRR"].fillna(0)
)
customer_summary["TotalSeats"] = (
    customer_summary["TotalSeats"].fillna(0)
)
customer_summary["SubscriptionCount"] = (
    customer_summary["SubscriptionCount"].fillna(0)
)
# STEP 7 — TENURE

today = pd.Timestamp.today().normalize()
customer_summary["TenureYears"] = (
    (
        today -
        customer_summary["SignupDate"]
    ).dt.days / 365.25
).round(2)

# STEP 8 — PIVOT SUMMARY
pivot_summary = pd.pivot_table(
    customer_summary,
    index="Industry",
    values=[
        "CustomerID",
        "TotalMRR",
        "TenureYears"
    ],
    aggfunc={
        "CustomerID": "nunique",
        "TotalMRR": "sum",
        "TenureYears": "mean"
    }
).reset_index()
pivot_summary = pivot_summary.rename(
    columns={
        "CustomerID": "CustomerCount",
        "TenureYears": "AverageTenureYears"
    }
)
pivot_summary["AverageMRRPerAccount"] = (
    pivot_summary["TotalMRR"] /
    pivot_summary["CustomerCount"]
)
pivot_summary[
    [
        "TotalMRR",
        "AverageTenureYears",
        "AverageMRRPerAccount"
    ]
] = pivot_summary[
    [
        "TotalMRR",
        "AverageTenureYears",
        "AverageMRRPerAccount"
    ]
].round(2)

print("\nPivot Summary:")
print(pivot_summary)

# STEP 9 — EXPORT TO EXCEL
output_file = "Module11_SaaS_Excel_Report.xlsx"
with pd.ExcelWriter(
    output_file,
    engine="openpyxl"
) as writer:
    customers.to_excel(
        writer,
        sheet_name="Customers",
        index=False
    )
    subscriptions.to_excel(
        writer,
        sheet_name="Subscriptions",
        index=False
    )
    usage.to_excel(
        writer,
        sheet_name="Usage",
        index=False
    )
    tickets.to_excel(
        writer,
        sheet_name="Tickets",
        index=False
    )
    customer_summary.to_excel(
        writer,
        sheet_name="Customer Summary",
        index=False
    )
    pivot_summary.to_excel(
        writer,
        sheet_name="Pivot Summary",
        index=False
    )
# STEP 10 — CREATE KPI SHEET

wb = load_workbook(output_file)
if "KPI" in wb.sheetnames:
    del wb["KPI"]
kpi = wb.create_sheet("KPI", 0)
customer_last_row = len(customer_summary) + 1


# Title
kpi["A1"] = "SaaS Analytics KPI Report"
kpi.merge_cells("A1:C1")

# Headers
kpi["A3"] = "KPI"
kpi["B3"] = "Value"
kpi["C3"] = "Formula / Definition"

# KPI 1 — TOTAL MRR
kpi["A4"] = "Total MRR"
kpi["B4"] = (
    f"=SUM('Customer Summary'!I2:I{customer_last_row})"
)
kpi["C4"] = "Sum of customer-level MRR"

# KPI 2 — CHURN RATE
kpi["A5"] = "Churn Rate"
kpi["B5"] = (
    f'=COUNTIF(\'Customer Summary\'!L2:L{customer_last_row},"Churned")'
    f'/COUNTIF(\'Customer Summary\'!L2:L{customer_last_row},"<>")'
)
kpi["C5"] = (
    "Churned customers / customers with latest subscription status"
)
# KPI 3 — AVERAGE REVENUE PER ACCOUNT

kpi["A6"] = "Average Revenue per Account"
kpi["B6"] = (
    f"=B4/COUNTA('Customer Summary'!A2:A{customer_last_row})"
)
kpi["C6"] = "Total MRR / customer accounts"

# KPI 4 — AVERAGE TENURE

kpi["A7"] = "Average Tenure"
kpi["B7"] = (
    f"=AVERAGE('Customer Summary'!M2:M{customer_last_row})"
)
kpi["C7"] = "Average customer tenure in years"

# NUMBER FORMATTING
kpi["B4"].number_format = '#,##0.00'
kpi["B5"].number_format = '0.00%'
kpi["B6"].number_format = '#,##0.00'
kpi["B7"].number_format = '0.00'

# FORMATTING

title_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)
header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)
kpi["A1"].fill = title_fill
kpi["A1"].font = Font(
    color="FFFFFF",
    bold=True,
    size=14
)
kpi["A1"].alignment = Alignment(
    horizontal="center"
)
for cell in kpi[3]:
    cell.fill = header_fill
    cell.font = Font(bold=True)
    cell.alignment = Alignment(
        horizontal="center"
    )
for row in range(4, 8):
    kpi[f"A{row}"].font = Font(bold=True)

kpi.column_dimensions["A"].width = 32
kpi.column_dimensions["B"].width = 22
kpi.column_dimensions["C"].width = 70

# FORMAT OTHER SHEETS

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    if sheet_name != "KPI":
        ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )
        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 2, 40)

# FORCE EXCEL TO RECALCULATE FORMULAS

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

# SAVE
wb.save(output_file)
print("\n======================================")
print("MODULE 11 COMPLETED")
print("======================================")
print(
    "Excel file created:",
    output_file
)

import pandas as pd

subscriptions = pd.read_csv(
    "cleaned_saas_subscriptions.csv"
)

python_total_mrr = subscriptions["MRR"].sum()

print("Python Total MRR:", round(python_total_mrr, 2))